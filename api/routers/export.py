"""
Glass Expert AI — Export Router
Export a conversation as a DOCX report.
"""
import io
from datetime import datetime
from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from loguru import logger
from api.auth import get_current_user, UserInToken
from api.database import get_db

router = APIRouter()


@router.get("/export/{session_id}")
async def export_conversation(
    session_id: str,
    user: UserInToken = Depends(get_current_user),
):
    """Export a conversation as a formatted DOCX report."""
    try:
        from docx import Document
        from docx.shared import Pt, Inches, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError:
        raise HTTPException(status_code=500, detail="python-docx not installed")

    # Fetch conversation messages
    with get_db() as conn:
        cur = conn.cursor()

        # Verify conversation belongs to user
        cur.execute(
            """SELECT DISTINCT session_id FROM chat_history
               WHERE session_id = %s AND user_id = %s""",
            (session_id, user.user_id),
        )
        if not cur.fetchone():
            raise HTTPException(status_code=404, detail="Conversation not found")

        # Get all messages
        cur.execute(
            """SELECT role, content, metadata, created_at
               FROM chat_history
               WHERE session_id = %s
               ORDER BY created_at ASC""",
            (session_id,),
        )
        messages = cur.fetchall()
        cur.close()

    if not messages:
        raise HTTPException(status_code=404, detail="No messages in conversation")

    # Build DOCX
    doc = Document()

    # Title
    title = doc.add_heading("Glass Expert AI — Consultation Report", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Metadata
    meta_para = doc.add_paragraph()
    meta_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = meta_para.add_run(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(128, 128, 128)

    doc.add_paragraph()  # spacer

    # Messages
    for role, content, metadata, created_at in messages:
        # Role header
        if role == "user":
            h = doc.add_heading("Question", level=2)
            for run in h.runs:
                run.font.color.rgb = RGBColor(30, 58, 95)
        else:
            h = doc.add_heading("Answer", level=2)
            for run in h.runs:
                run.font.color.rgb = RGBColor(20, 80, 50)

        # Timestamp
        ts_para = doc.add_paragraph()
        ts_run = ts_para.add_run(
            created_at.strftime("%Y-%m-%d %H:%M") if hasattr(created_at, 'strftime') else str(created_at)
        )
        ts_run.font.size = Pt(8)
        ts_run.font.color.rgb = RGBColor(160, 160, 160)

        # Content
        content_para = doc.add_paragraph(content or "")
        content_para.style.font.size = Pt(11)

        # Sources (if assistant message with metadata)
        if role == "assistant" and metadata:
            meta = metadata if isinstance(metadata, dict) else {}
            sources = meta.get("sources", [])
            if sources:
                doc.add_heading("Sources", level=3)
                for i, src in enumerate(sources, 1):
                    src_title = src.get("title", "Unknown")
                    src_type = src.get("source_type", "")
                    doc.add_paragraph(
                        f"[{i}] {src_title} ({src_type})",
                        style="List Number",
                    )

        doc.add_paragraph()  # spacer between Q&A pairs

    # Footer
    doc.add_paragraph()
    footer = doc.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Glass Expert AI · RAG-powered Glass Science Assistant")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(160, 160, 160)

    # Stream as download
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    filename = f"glass_expert_report_{session_id[:8]}_{datetime.now().strftime('%Y%m%d')}.docx"

    logger.info(f"Export: conversation {session_id[:8]} by {user.email}")

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
